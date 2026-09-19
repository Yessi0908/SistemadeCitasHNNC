from datetime import date, timedelta
from calendar import monthrange
from pathlib import Path
from django.conf import settings
from django.core.exceptions import ValidationError as ErrorValidacionDjango
from django.utils import timezone
from django.contrib.auth import get_user_model
from django.db.models import Count, Q
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework.filters import SearchFilter
from rest_framework import viewsets, status, generics
from rest_framework.decorators import action, api_view, permission_classes
from rest_framework.response import Response
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework_simplejwt.tokens import RefreshToken
from rest_framework.views import APIView
from rest_framework.parsers import JSONParser, FormParser
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from django.http import HttpResponse

from .models import (
    Paciente, Consulta, Cita, Bitacora, ExpedienteVAS, NotaVAS,
    EstadisticaRegistro, RegistroDiario, RegistroDiarioDetalle,
    TokenListaNegra, Rol, Medico,
)
from .serializers import (
    PacienteSerializer, ConsultaSerializer, CitaSerializer, BitacoraSerializer,
    VASSerializer, NotaVASSerializer, EstadisticaSerializer, UsuarioSerializer, UsuarioCrearSerializer,
    RegistroDiarioSerializer, RegistroDiarioDetalleSerializer, MedicoSerializer,
    UsuarioDetalleAdminSerializer,
)
from .permissions import (
    EsAdmin, EsArchivo, EsRegistros, EsEstadistica, EsJuridico, SoloLecturaArchivo,
    EsLecturaArchivoRegistros,
)
from .utils import (
    registrar_bitacora, generar_numero_expediente, usuario_bloqueado, obtener_ip,
)
from . import pdf as generador_pdf

Usuario = get_user_model()


def justificacion_requerida(request, minimo=8):
    texto = ''
    if hasattr(request, 'data'):
        texto = str(request.data.get('justificacion') or '').strip()
    if not texto:
        texto = str(request.query_params.get('justificacion') or '').strip()
    if len(texto) < minimo:
        return None
    return texto


class LoginView(APIView):
    permission_classes = [AllowAny]
    parser_classes = [JSONParser, FormParser]

    def post(self, request):
        username = request.data.get('username', '').strip()
        password = request.data.get('password', '')
        ip = obtener_ip(request)

        try:
            user = Usuario.objects.get(username=username)
        except Usuario.DoesNotExist:
            registrar_bitacora(username, 'LOGIN_FALLIDO', 'Usuario no existe', ip)
            return Response({'error': 'Credenciales inválidas'}, status=401)

        if not user.activo_sistema or not user.is_active:
            return Response({'error': 'Usuario inactivo'}, status=403)

        if usuario_bloqueado(user):
            return Response({'error': 'Cuenta bloqueada temporalmente'}, status=403)

        if not user.check_password(password):
            user.intentos_fallidos += 1
            from django.conf import settings
            if user.intentos_fallidos >= settings.MAX_INTENTOS_LOGIN:
                user.bloqueado_hasta = timezone.now() + timedelta(minutes=settings.BLOQUEO_MINUTOS)
                user.intentos_fallidos = 0
            user.save(update_fields=['intentos_fallidos', 'bloqueado_hasta'])
            registrar_bitacora(username, 'LOGIN_FALLIDO', f'Intento {user.intentos_fallidos}', ip, user.rol)
            return Response({'error': 'Credenciales inválidas'}, status=401)

        user.intentos_fallidos = 0
        user.bloqueado_hasta = None
        user.save(update_fields=['intentos_fallidos', 'bloqueado_hasta'])

        refresh = RefreshToken.for_user(user)
        refresh['rol'] = user.rol
        refresh['username'] = user.username
        access = refresh.access_token
        access['rol'] = user.rol

        registrar_bitacora(username, 'LOGIN_OK', '', ip, user.rol)
        return Response({
            'access': str(access),
            'refresh': str(refresh),
            'rol': user.rol,
            'username': user.username,
            'user_id': user.id,
            'nombre': user.get_full_name() or user.username,
        })


class LogoutView(APIView):
    parser_classes = [JSONParser, FormParser]

    def post(self, request):
        token = request.data.get('refresh')
        if token:
            try:
                rt = RefreshToken(token)
                TokenListaNegra.objects.get_or_create(jti=str(rt['jti']), usuario=request.user.username)
                rt.blacklist()
            except Exception:
                pass
        registrar_bitacora(request.user.username, 'LOGOUT', '', obtener_ip(request), request.user.rol)
        return Response({'mensaje': 'Sesión cerrada'})


class PacienteViewSet(viewsets.ModelViewSet):
    queryset = Paciente.objects.all()
    serializer_class = PacienteSerializer
    permission_classes = [IsAuthenticated, SoloLecturaArchivo]
    filter_backends = [DjangoFilterBackend, SearchFilter]
    filterset_fields = ['especialidad', 'estado_paciente']
    search_fields = ['primer_nombre', 'segundo_nombre', 'primer_apellido', 'segundo_apellido', 'dpi', 'numero_expediente']

    def get_permissions(self):
        if self.action in ('create', 'update', 'partial_update'):
            return [IsAuthenticated(), EsRegistros()]
        if self.action == 'destroy':
            return [IsAuthenticated(), EsRegistros()]
        return super().get_permissions()

    def perform_create(self, serializer):
        expediente = generar_numero_expediente()
        paciente = serializer.save(numero_expediente=expediente)
        registrar_bitacora(
            self.request.user.username, 'CREAR_PACIENTE',
            paciente.numero_expediente, obtener_ip(self.request), self.request.user.rol,
        )

    def perform_update(self, serializer):
        anterior = serializer.instance.estado_paciente
        paciente = serializer.save()
        if anterior != paciente.estado_paciente:
            registrar_bitacora(
                self.request.user.username, 'CAMBIAR_ESTADO_PACIENTE',
                f'{paciente.numero_expediente}: {anterior} → {paciente.estado_paciente}',
                obtener_ip(self.request), self.request.user.rol,
            )
        else:
            registrar_bitacora(
                self.request.user.username, 'ACTUALIZAR_PACIENTE',
                paciente.numero_expediente, obtener_ip(self.request), self.request.user.rol,
            )

    def perform_destroy(self, instance):
        expediente = instance.numero_expediente
        instance.delete()
        registrar_bitacora(
            self.request.user.username, 'ELIMINAR_PACIENTE',
            expediente, obtener_ip(self.request), self.request.user.rol,
        )

    @action(detail=True, methods=['get'], permission_classes=[IsAuthenticated])
    def carnet(self, request, pk=None):
        if request.user.rol not in (Rol.REGISTROS, Rol.ADMIN, Rol.ARCHIVO):
            return Response({'error': 'Sin permiso'}, status=403)
        return generador_pdf.generar_carnet(self.get_object())

    @action(detail=True, methods=['get'])
    def hoja_expediente(self, request, pk=None):
        paciente = self.get_object()
        consultas = paciente.consultas.all()
        return generador_pdf.generar_hoja_expediente(paciente, consultas)

    @action(detail=True, methods=['get'])
    def constancia(self, request, pk=None):
        return generador_pdf.generar_constancia_laboral(self.get_object())

    @action(detail=True, methods=['get'], permission_classes=[IsAuthenticated, EsLecturaArchivoRegistros])
    def citas_resumen(self, request, pk=None):
        paciente = self.get_object()
        hoy = date.today()
        citas = Cita.objects.filter(paciente=paciente).order_by('fecha', 'hora')
        proximas = citas.filter(
            fecha__gte=hoy, estado='Confirmada'
        ).order_by('fecha', 'hora')
        historial = citas.exclude(
            pk__in=proximas.values_list('pk', flat=True)
        ).order_by('-fecha', '-hora')
        ser = CitaSerializer
        return Response({
            'paciente_id': paciente.id,
            'numero_expediente': paciente.numero_expediente,
            'nombre': paciente.nombre_completo,
            'edad_texto': paciente.edad_texto,
            'proximas': ser(proximas, many=True).data,
            'historial': ser(historial, many=True).data,
        })


class ConsultaViewSet(viewsets.ModelViewSet):
    queryset = Consulta.objects.select_related('paciente')
    serializer_class = ConsultaSerializer
    permission_classes = [IsAuthenticated, SoloLecturaArchivo]
    filterset_fields = ['paciente', 'fecha', 'especialidad']

    def get_permissions(self):
        if self.request.method not in ('GET', 'HEAD', 'OPTIONS'):
            return [IsAuthenticated(), EsRegistros()]
        if getattr(self.request.user, 'is_authenticated', False) and self.request.user.rol == Rol.ARCHIVO:
            return [IsAuthenticated(), EsArchivo()]
        return [IsAuthenticated()]

    def perform_create(self, serializer):
        serializer.save(creado_por=self.request.user.username)


class MedicoViewSet(viewsets.ModelViewSet):
    queryset = Medico.objects.all()
    serializer_class = MedicoSerializer
    permission_classes = [IsAuthenticated]
    filterset_fields = ['especialidad', 'activo']

    def get_permissions(self):
        if self.action in ('create', 'update', 'partial_update', 'destroy'):
            return [IsAuthenticated(), EsAdmin()]
        return [IsAuthenticated()]

    def get_queryset(self):
        qs = Medico.objects.all()
        if self.request.query_params.get('activos') == '1':
            qs = qs.filter(activo=True)
        esp = self.request.query_params.get('especialidad', '').strip()
        if esp:
            qs = qs.filter(especialidad=esp)
        return qs.order_by('especialidad', 'nombre')

    def perform_create(self, serializer):
        medico = serializer.save()
        registrar_bitacora(
            self.request.user.username, 'CREAR_MEDICO',
            f'{medico.nombre} — {medico.especialidad}', obtener_ip(self.request), self.request.user.rol,
        )

    def _parse_bool(self, valor):
        if isinstance(valor, bool):
            return valor
        if valor is None:
            return None
        return str(valor).strip().lower() in ('1', 'true', 't', 'si', 'sí', 'yes')

    def partial_update(self, request, *args, **kwargs):
        instance = self.get_object()
        if 'activo' in request.data:
            nuevo = self._parse_bool(request.data.get('activo'))
            if nuevo is not None and bool(nuevo) != bool(instance.activo):
                motivo = justificacion_requerida(request)
                if not motivo:
                    return Response({
                        'error': 'Debe indicar una justificación para cambiar el estado del médico.',
                    }, status=400)
                self._justificacion = motivo
                self._cambio_activo = 'ACTIVAR_MEDICO' if nuevo else 'DESACTIVAR_MEDICO'
        return super().partial_update(request, *args, **kwargs)

    def perform_update(self, serializer):
        medico = serializer.save()
        if getattr(self, '_cambio_activo', None):
            registrar_bitacora(
                self.request.user.username, self._cambio_activo,
                f'{medico.nombre} — {medico.especialidad}. Justificación: {getattr(self, "_justificacion", "")}',
                obtener_ip(self.request), self.request.user.rol,
            )
        else:
            registrar_bitacora(
                self.request.user.username, 'ACTUALIZAR_MEDICO',
                f'{medico.nombre} — {medico.especialidad}', obtener_ip(self.request), self.request.user.rol,
            )

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        motivo = justificacion_requerida(request)
        if not motivo:
            return Response({
                'error': 'Debe indicar una justificación para eliminar al médico.',
            }, status=400)
        citas_abiertas = instance.citas.exclude(estado__in=('Atendida', 'Cancelada'))
        if citas_abiertas.exists():
            return Response({
                'error': (
                    'No se puede eliminar el médico porque tiene citas confirmadas o pendientes. '
                    'Cuando todas estén atendidas o canceladas sí se podrá eliminar.'
                ),
            }, status=400)
        instance.citas.update(medico_nombre=instance.nombre, medico=None)
        self._justificacion = motivo
        return super().destroy(request, *args, **kwargs)

    def perform_destroy(self, instance):
        detalle = (
            f'{instance.nombre} — {instance.especialidad}. '
            f'Justificación: {getattr(self, "_justificacion", "")}'
        )
        instance.delete()
        registrar_bitacora(
            self.request.user.username, 'ELIMINAR_MEDICO', detalle, obtener_ip(self.request), self.request.user.rol,
        )


class CitaViewSet(viewsets.ModelViewSet):
    queryset = Cita.objects.select_related('paciente', 'medico')
    serializer_class = CitaSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend, SearchFilter]
    filterset_fields = ['fecha', 'estado', 'especialidad', 'paciente']
    search_fields = ['paciente__numero_expediente', 'paciente__dpi', 'paciente__primer_nombre', 'paciente__primer_apellido']

    def get_permissions(self):
        if self.request.method in ('GET', 'HEAD', 'OPTIONS'):
            return [IsAuthenticated(), EsLecturaArchivoRegistros()]
        return [IsAuthenticated(), EsRegistros()]

    def perform_create(self, serializer):
        cita = serializer.save()
        registrar_bitacora(
            self.request.user.username, 'CREAR_CITA',
            f'{cita.paciente.numero_expediente} {cita.fecha}', obtener_ip(self.request), self.request.user.rol,
        )

    def perform_update(self, serializer):
        cita = serializer.save()
        registrar_bitacora(
            self.request.user.username, 'ACTUALIZAR_CITA',
            f'{cita.paciente.numero_expediente} {cita.fecha}', obtener_ip(self.request), self.request.user.rol,
        )

    def perform_destroy(self, instance):
        detalle = f'{instance.paciente.numero_expediente} {instance.fecha}'
        instance.delete()
        registrar_bitacora(
            self.request.user.username, 'ELIMINAR_CITA', detalle, obtener_ip(self.request), self.request.user.rol,
        )

    @action(detail=False, methods=['get'])
    def alertas(self, request):
        manana = date.today() + timedelta(days=1)
        citas = self.get_queryset().filter(fecha=manana, estado='Confirmada')[:20]
        return Response(CitaSerializer(citas, many=True).data)


class VASViewSet(viewsets.ModelViewSet):
    queryset = ExpedienteVAS.objects.select_related('paciente')
    serializer_class = VASSerializer
    permission_classes = [IsAuthenticated, EsJuridico]
    search_fields = ['paciente__dpi', 'paciente__numero_expediente']
    ESP_VAS = 'Violencia y Abuso Sexual (VAS)'

    def perform_create(self, serializer):
        serializer.save(creado_por=self.request.user.username)

    def _paciente_vas(self, paciente_id):
        paciente = Paciente.objects.filter(pk=paciente_id, especialidad=self.ESP_VAS).first()
        if not paciente:
            return None
        return paciente

    def _detalle_paciente(self, paciente):
        hoy = date.today()
        citas = Cita.objects.filter(paciente=paciente).order_by('fecha', 'hora')
        proximas = citas.filter(fecha__gte=hoy, estado='Confirmada').order_by('fecha', 'hora')
        historial = citas.exclude(pk__in=proximas.values_list('pk', flat=True)).order_by('-fecha', '-hora')
        notas = paciente.notas_vas.all()
        return {
            'paciente_id': paciente.id,
            'numero_expediente': paciente.numero_expediente,
            'nombre': paciente.nombre_completo,
            'dpi': paciente.dpi,
            'estado_paciente': paciente.estado_paciente,
            'proximas': CitaSerializer(proximas, many=True).data,
            'historial': CitaSerializer(historial, many=True).data,
            'notas': NotaVASSerializer(notas, many=True).data,
        }

    @action(detail=False, methods=['get'])
    def pacientes(self, request):
        qs = Paciente.objects.filter(especialidad=self.ESP_VAS)
        q = request.query_params.get('q', '').strip()
        if q:
            qs = qs.filter(
                Q(numero_expediente__icontains=q) |
                Q(dpi__icontains=q) |
                Q(primer_apellido__icontains=q) |
                Q(primer_nombre__icontains=q)
            )
        qs = qs.order_by('primer_apellido', 'primer_nombre')[:200]
        return Response(PacienteSerializer(qs, many=True).data)

    @action(detail=False, methods=['get'], url_path=r'paciente/(?P<paciente_id>[0-9]+)/detalle')
    def detalle(self, request, paciente_id=None):
        paciente = self._paciente_vas(paciente_id)
        if not paciente:
            return Response({'error': 'Paciente VAS no encontrado.'}, status=404)
        return Response(self._detalle_paciente(paciente))

    @action(detail=False, methods=['post'], url_path=r'paciente/(?P<paciente_id>[0-9]+)/notas')
    def crear_nota(self, request, paciente_id=None):
        paciente = self._paciente_vas(paciente_id)
        if not paciente:
            return Response({'error': 'Paciente VAS no encontrado.'}, status=404)
        texto = str(request.data.get('texto') or '').strip()
        if len(texto) < 5:
            return Response({'error': 'La nota debe tener al menos 5 caracteres.'}, status=400)
        nota = NotaVAS.objects.create(
            paciente=paciente,
            texto=texto,
            creado_por=request.user.username,
        )
        registrar_bitacora(
            request.user.username, 'NOTA_JURIDICO',
            f'{paciente.numero_expediente}: {texto}',
            obtener_ip(request), request.user.rol,
        )
        return Response(self._detalle_paciente(paciente), status=201)

    @action(detail=False, methods=['post'], url_path=r'notas/(?P<nota_id>[0-9]+)/confirmar')
    def confirmar_nota(self, request, nota_id=None):
        nota = NotaVAS.objects.select_related('paciente').filter(pk=nota_id).first()
        if not nota or nota.paciente.especialidad != self.ESP_VAS:
            return Response({'error': 'Nota no encontrada.'}, status=404)
        if nota.lista:
            return Response({
                'error': 'Esta nota ya está confirmada y no se puede modificar, ni siquiera el administrador.',
            }, status=403)
        nota.lista = True
        nota.confirmada_por = request.user.username
        nota.confirmada = timezone.now()
        try:
            nota.save(update_fields=['lista', 'confirmada_por', 'confirmada'])
        except ErrorValidacionDjango as e:
            return Response({'error': str(e)}, status=403)
        registrar_bitacora(
            request.user.username, 'NOTA_JURIDICO_LISTA',
            f'{nota.paciente.numero_expediente}: nota confirmada como lista',
            obtener_ip(request), request.user.rol,
        )
        return Response(self._detalle_paciente(nota.paciente))

    @action(detail=False, methods=['post'], url_path=r'notas/(?P<nota_id>[0-9]+)/editar')
    def editar_nota(self, request, nota_id=None):
        nota = NotaVAS.objects.select_related('paciente').filter(pk=nota_id).first()
        if not nota or nota.paciente.especialidad != self.ESP_VAS:
            return Response({'error': 'Nota no encontrada.'}, status=404)
        if nota.lista:
            return Response({
                'error': 'Esta nota ya está confirmada. No se puede modificar, ni siquiera el administrador.',
            }, status=403)
        texto = str(request.data.get('texto') or '').strip()
        if len(texto) < 5:
            return Response({'error': 'La nota debe tener al menos 5 caracteres.'}, status=400)
        nota.texto = texto
        try:
            nota.save(update_fields=['texto'])
        except ErrorValidacionDjango as e:
            return Response({'error': str(e)}, status=403)
        registrar_bitacora(
            request.user.username, 'NOTA_JURIDICO',
            f'{nota.paciente.numero_expediente} (borrador editado): {texto}',
            obtener_ip(request), request.user.rol,
        )
        return Response(self._detalle_paciente(nota.paciente))


class EstadisticaViewSet(viewsets.ModelViewSet):
    queryset = EstadisticaRegistro.objects.all()
    serializer_class = EstadisticaSerializer
    permission_classes = [IsAuthenticated, EsEstadistica]

    @staticmethod
    def _periodo_mes(mes_iso=None):
        hoy = date.today()
        anio, mes = hoy.year, hoy.month
        if mes_iso:
            try:
                partes = str(mes_iso).strip()[:7].split('-')
                anio, mes = int(partes[0]), int(partes[1])
                date(anio, mes, 1)
            except (ValueError, TypeError, IndexError):
                anio, mes = hoy.year, hoy.month
        inicio = date(anio, mes, 1)
        fin = date(anio, mes, monthrange(anio, mes)[1])
        return inicio, fin, f'{anio:04d}-{mes:02d}'

    @classmethod
    def _datos_resumen(cls, mes_iso=None):
        inicio, fin, mes_id = cls._periodo_mes(mes_iso)
        return {
            'mes': mes_id,
            'pacientes_por_especialidad': list(
                Paciente.objects.values('especialidad').annotate(total=Count('id')).order_by('-total')[:50]
            ),
            'pacientes_por_estado': list(
                Paciente.objects.values('estado_paciente').annotate(total=Count('id'))
            ),
            'consultas_mes': list(
                Consulta.objects.filter(fecha__gte=inicio, fecha__lte=fin)
                .values('especialidad').annotate(total=Count('id'))
            ),
            'citas_por_especialidad': list(
                Cita.objects.values('especialidad').annotate(total=Count('id')).order_by('-total')
            ),
            'citas_por_estado': list(
                Cita.objects.values('estado').annotate(total=Count('id'))
            ),
            'citas_mes': list(
                Cita.objects.filter(fecha__gte=inicio, fecha__lte=fin)
                .values('especialidad').annotate(total=Count('id')).order_by('-total')
            ),
        }

    @action(detail=False, methods=['get'])
    def resumen_tablas(self, request):
        return Response(self._datos_resumen(request.query_params.get('mes')))

    @action(detail=False, methods=['get'])
    def exportar_excel(self, request):
        datos = self._datos_resumen(request.query_params.get('mes'))
        wb = Workbook()
        ws = wb.active
        ws.title = 'Estadísticas HNNCJ'
        azul = Font(name='Calibri', bold=True, color='003366', size=14)
        azul_sub = Font(name='Calibri', bold=True, color='003366', size=11)
        titulo_sec = Font(name='Calibri', bold=True, color='003366', size=12)
        encabezado = Font(name='Calibri', bold=True, color='FFFFFF')
        relleno = PatternFill('solid', fgColor='003366')

        ws.merge_cells('B1:E1')
        ws.merge_cells('B2:E2')
        ws.merge_cells('B3:E3')
        ws['B1'] = 'MINISTERIO DE SALUD PÚBLICA Y ASISTENCIA SOCIAL'
        ws['B1'].font = azul
        ws['B2'] = 'Hospital Nacional Nicolasa Cruz Jalapa'
        ws['B2'].font = azul
        ws['B3'] = (
            f'Reporte de estadísticas — {date.today().strftime("%d/%m/%Y")} '
            f'(citas/consultas del mes {datos.get("mes") or ""})'
        )
        ws['B3'].font = azul_sub
        for celda in ('B1', 'B2', 'B3'):
            ws[celda].alignment = Alignment(horizontal='left', vertical='center')

        logo = Path(getattr(settings, 'LOGO_RUTA', '') or '')
        if logo.exists():
            try:
                img = XLImage(str(logo))
                img.width = 72
                img.height = 72
                ws.add_image(img, 'A1')
            except Exception:
                pass
        ws.row_dimensions[1].height = 22
        ws.row_dimensions[2].height = 20
        ws.row_dimensions[3].height = 18
        ws.column_dimensions['A'].width = 42
        ws.column_dimensions['B'].width = 28
        ws.column_dimensions['C'].width = 16
        ws.column_dimensions['D'].width = 16

        fila = 5

        def seccion(titulo, encabezados, filas):
            nonlocal fila
            ws.cell(fila, 1, titulo).font = titulo_sec
            fila += 1
            for i, h in enumerate(encabezados, 1):
                celda = ws.cell(fila, i, h)
                celda.font = encabezado
                celda.fill = relleno
            fila += 1
            if not filas:
                ws.cell(fila, 1, 'Sin datos')
                fila += 1
            else:
                for valores in filas:
                    for i, valor in enumerate(valores, 1):
                        ws.cell(fila, i, valor)
                    fila += 1
            fila += 2

        seccion(
            'Pacientes por especialidad',
            ['Especialidad', 'Total'],
            [(r.get('especialidad') or '—', r.get('total') or 0) for r in datos['pacientes_por_especialidad']],
        )
        seccion(
            'Pacientes por estado',
            ['Estado', 'Total'],
            [(r.get('estado_paciente') or '—', r.get('total') or 0) for r in datos['pacientes_por_estado']],
        )
        seccion(
            f'Consultas del mes {datos.get("mes") or ""}',
            ['Especialidad', 'Total'],
            [(r.get('especialidad') or '—', r.get('total') or 0) for r in datos['consultas_mes']],
        )
        seccion(
            'Citas por especialidad',
            ['Especialidad', 'Total'],
            [(r.get('especialidad') or '—', r.get('total') or 0) for r in datos['citas_por_especialidad']],
        )
        seccion(
            'Citas por estado',
            ['Estado', 'Total'],
            [(r.get('estado') or '—', r.get('total') or 0) for r in datos['citas_por_estado']],
        )
        seccion(
            f'Citas del mes {datos.get("mes") or ""}',
            ['Especialidad', 'Total'],
            [(r.get('especialidad') or '—', r.get('total') or 0) for r in datos['citas_mes']],
        )
        extras = EstadisticaRegistro.objects.all()[:500]
        if extras:
            seccion(
                'Registros manuales',
                ['Categoría', 'Valor', 'Cantidad', 'Fecha'],
                [(e.categoria, e.valor, e.cantidad, str(e.fecha)) for e in extras],
            )
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=estadisticas.xlsx'
        wb.save(response)
        return response


class BitacoraViewSet(viewsets.ReadOnlyModelViewSet):
    serializer_class = BitacoraSerializer
    permission_classes = [IsAuthenticated, EsAdmin]
    queryset = Bitacora.objects.all()
    pagination_class = None

    def get_queryset(self):
        qs = Bitacora.objects.all()
        accion = self.request.query_params.get('accion', '').strip()
        usuario = self.request.query_params.get('usuario', '').strip()
        q = self.request.query_params.get('q', '').strip()
        if accion == 'NOTAS_JURIDICO':
            qs = qs.filter(accion__in=('NOTA_JURIDICO', 'NOTA_JURIDICO_LISTA'))
        elif accion:
            qs = qs.filter(accion=accion)
        if usuario:
            qs = qs.filter(usuario__icontains=usuario)
        if q:
            qs = qs.filter(
                Q(accion__icontains=q) | Q(detalle__icontains=q) | Q(usuario__icontains=q)
            )
        return qs[:300]

    @action(detail=False, methods=['get'])
    def filtros(self, request):
        return Response({
            'acciones': [
                {'id': 'CREAR_PACIENTE', 'texto': 'Creación de pacientes'},
                {'id': 'ACTUALIZAR_PACIENTE', 'texto': 'Actualización de pacientes'},
                {'id': 'CAMBIAR_ESTADO_PACIENTE', 'texto': 'Cambio de estado del paciente'},
                {'id': 'ELIMINAR_PACIENTE', 'texto': 'Eliminación de pacientes'},
                {'id': 'CREAR_CITA', 'texto': 'Creación de citas'},
                {'id': 'ACTUALIZAR_CITA', 'texto': 'Actualización de citas'},
                {'id': 'ELIMINAR_CITA', 'texto': 'Eliminación de citas'},
                {'id': 'CREAR_MEDICO', 'texto': 'Creación de médicos'},
                {'id': 'ACTIVAR_MEDICO', 'texto': 'Activación de médicos'},
                {'id': 'DESACTIVAR_MEDICO', 'texto': 'Desactivación de médicos'},
                {'id': 'ELIMINAR_MEDICO', 'texto': 'Eliminación de médicos'},
                {'id': 'CREAR_USUARIO', 'texto': 'Creación de usuarios'},
                {'id': 'ELIMINAR_USUARIO', 'texto': 'Eliminación de usuarios'},
                {'id': 'BLOQUEAR_USUARIO', 'texto': 'Bloqueo de usuarios'},
                {'id': 'REACTIVAR_USUARIO', 'texto': 'Reactivación de usuarios'},
                {'id': 'NOTAS_JURIDICO', 'texto': 'Notas de jurídico'},
                {'id': 'NOTA_JURIDICO', 'texto': 'Nota jurídica agregada'},
                {'id': 'NOTA_JURIDICO_LISTA', 'texto': 'Nota jurídica confirmada'},
                {'id': 'LOGIN_OK', 'texto': 'Inicios de sesión'},
                {'id': 'LOGIN_FALLIDO', 'texto': 'Intentos fallidos de acceso'},
                {'id': 'LOGOUT', 'texto': 'Cierres de sesión'},
            ]
        })


class UsuarioViewSet(viewsets.ModelViewSet):
    queryset = Usuario.objects.all()
    permission_classes = [IsAuthenticated, EsAdmin]

    def get_serializer_class(self):
        if self.action == 'create':
            return UsuarioCrearSerializer
        return UsuarioSerializer

    def perform_create(self, serializer):
        user = serializer.save()
        registrar_bitacora(self.request.user.username, 'CREAR_USUARIO', user.username, obtener_ip(self.request))

    def perform_destroy(self, instance):
        username = instance.username
        instance.delete()
        detalle = f'{username}. Justificación: {getattr(self, "_justificacion", "")}'
        registrar_bitacora(
            self.request.user.username, 'ELIMINAR_USUARIO', detalle, obtener_ip(self.request),
        )

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        if instance.pk == request.user.pk:
            return Response({'error': 'No puede eliminar su propia cuenta.'}, status=400)
        motivo = justificacion_requerida(request)
        if not motivo:
            return Response({
                'error': 'Debe indicar una justificación para eliminar al usuario.',
            }, status=400)
        self._justificacion = motivo
        self.perform_destroy(instance)
        return Response(status=status.HTTP_204_NO_CONTENT)

    @action(detail=True, methods=['post'])
    def bloquear(self, request, pk=None):
        usuario = self.get_object()
        if usuario.pk == request.user.pk:
            return Response({'error': 'No puede bloquear su propia cuenta.'}, status=400)
        usuario.is_active = False
        usuario.activo_sistema = False
        usuario.save(update_fields=['is_active', 'activo_sistema'])
        registrar_bitacora(
            request.user.username, 'BLOQUEAR_USUARIO', usuario.username, obtener_ip(request),
        )
        return Response(UsuarioSerializer(usuario).data)

    @action(detail=True, methods=['post'])
    def reactivar(self, request, pk=None):
        usuario = self.get_object()
        usuario.is_active = True
        usuario.activo_sistema = True
        usuario.intentos_fallidos = 0
        usuario.bloqueado_hasta = None
        usuario.save(update_fields=['is_active', 'activo_sistema', 'intentos_fallidos', 'bloqueado_hasta'])
        registrar_bitacora(
            request.user.username, 'REACTIVAR_USUARIO', usuario.username, obtener_ip(request),
        )
        return Response(UsuarioSerializer(usuario).data)

    @action(detail=True, methods=['post'])
    def ver_datos(self, request, pk=None):
        usuario = self.get_object()
        clave_admin = request.data.get('contrasena_admin', '')
        if not clave_admin:
            return Response({'error': 'Debe ingresar su contraseña de administrador.'}, status=400)
        if not request.user.check_password(clave_admin):
            return Response({'error': 'Contraseña de administrador incorrecta.'}, status=403)
        registrar_bitacora(
            request.user.username, 'VER_DATOS_USUARIO', usuario.username, obtener_ip(request),
        )
        return Response(UsuarioDetalleAdminSerializer(usuario).data)


class RegistroDiarioView(APIView):
    permission_classes = [IsAuthenticated]

    def get_permissions(self):
        if self.request.method == 'POST':
            return [IsAuthenticated(), EsRegistros()]
        return [IsAuthenticated(), EsLecturaArchivoRegistros()]

    @staticmethod
    def _parse_fecha(fecha_str):
        try:
            f = date.fromisoformat(str(fecha_str)[:10])
        except (TypeError, ValueError):
            raise ValueError('Fecha no válida')
        if f > date.today():
            raise ValueError('La fecha del registro no puede ser futura.')
        return f

    def get(self, request):
        try:
            fecha = self._parse_fecha(request.query_params.get('fecha', str(date.today())))
        except ValueError as e:
            return Response({'error': str(e)}, status=400)
        registro, _ = RegistroDiario.objects.get_or_create(fecha=fecha)
        detalles = registro.detalles.select_related('paciente')
        return Response({
            'registro': RegistroDiarioSerializer(registro).data,
            'detalles': RegistroDiarioDetalleSerializer(detalles, many=True).data,
        })

    def post(self, request):
        try:
            fecha = self._parse_fecha(request.data.get('fecha', str(date.today())))
        except ValueError as e:
            return Response({'error': str(e)}, status=400)
        paciente_id = request.data.get('paciente_id')
        especialidad = request.data.get('especialidad', '')
        registro, _ = RegistroDiario.objects.get_or_create(fecha=fecha)
        paciente = Paciente.objects.get(pk=paciente_id)
        orden = registro.detalles.count() + 1
        RegistroDiarioDetalle.objects.create(
            registro=registro, paciente=paciente, orden=orden, especialidad=especialidad,
        )
        registro.total_pacientes = registro.detalles.count()
        registro.save()
        return Response({'mensaje': 'Agregado al registro diario'})

    def put(self, request):
        try:
            fecha = self._parse_fecha(request.query_params.get('fecha', str(date.today())))
        except ValueError as e:
            return Response({'error': str(e)}, status=400)
        registro = RegistroDiario.objects.filter(fecha=fecha).first()
        if not registro:
            return Response({'error': 'Sin registro'}, status=404)
        detalles = registro.detalles.select_related('paciente')
        return generador_pdf.generar_registro_diario(registro, detalles)


class RespaldoView(APIView):
    permission_classes = [IsAuthenticated, EsAdmin]

    def post(self, request):
        import json
        from django.core import serializers as ser
        datos = {
            'pacientes': json.loads(ser.serialize('json', Paciente.objects.all())),
            'consultas': json.loads(ser.serialize('json', Consulta.objects.all())),
        }
        registrar_bitacora(request.user.username, 'RESPALDO', 'JSON manual', obtener_ip(request))
        return Response({'mensaje': 'Respaldo generado', 'registros': len(datos['pacientes'])})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def catalogos_paciente(request):
    from .catalogos import (
        ESPECIALIDADES, ESTADOS_PACIENTE, ESTADOS_CIVILES, ESTADOS_CITA,
        ESPECIALIDADES_TARDE,
    )
    medicos = Medico.objects.filter(activo=True).order_by('especialidad', 'nombre')
    return Response({
        'especialidades': ESPECIALIDADES,
        'especialidades_tarde': ESPECIALIDADES_TARDE,
        'estados_paciente': ESTADOS_PACIENTE,
        'estados_civiles': ESTADOS_CIVILES,
        'estados_cita': ESTADOS_CITA,
        'medicos': MedicoSerializer(medicos, many=True).data,
        'fecha_hoy': str(date.today()),
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def mi_perfil(request):
    return Response({
        'username': request.user.username,
        'rol': request.user.rol,
        'nombre': request.user.get_full_name(),
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated, EsRegistros])
def busqueda_registros(request):
    q = request.query_params.get('q', '').strip()
    if not q:
        return Response([])
    pacientes = Paciente.objects.filter(
        Q(dpi__icontains=q) |
        Q(numero_expediente__icontains=q) |
        Q(primer_nombre__icontains=q) |
        Q(segundo_nombre__icontains=q) |
        Q(primer_apellido__icontains=q) |
        Q(segundo_apellido__icontains=q)
    )[:30]
    return Response(PacienteSerializer(pacientes, many=True).data)


@api_view(['GET'])
@permission_classes([IsAuthenticated, EsArchivo])
def archivo_buscar_citas(request):
    fecha = request.query_params.get('fecha', '').strip()
    especialidad = request.query_params.get('especialidad', '').strip()
    citas = Cita.objects.select_related('paciente', 'medico').all()
    if fecha:
        citas = citas.filter(fecha=fecha)
    if especialidad:
        citas = citas.filter(especialidad=especialidad)
    citas = citas.order_by('fecha', 'hora')[:200]
    return Response(CitaSerializer(citas, many=True).data)


@api_view(['GET'])
@permission_classes([IsAuthenticated, EsArchivo])
def busqueda_archivo(request):
    q = request.query_params.get('q', '').strip()
    if not q:
        return Response({'resultados': []})
    pacientes = Paciente.objects.filter(
        Q(dpi__icontains=q) |
        Q(numero_expediente__icontains=q) |
        Q(primer_nombre__icontains=q) |
        Q(primer_apellido__icontains=q)
    )[:30]
    return Response(PacienteSerializer(pacientes, many=True).data)
